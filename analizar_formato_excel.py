import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image
import os

# Cargar el archivo
wb = openpyxl.load_workbook('C:/analisis-produccion-stc/Resumen_Dia.xlsx')
ws = wb.active

print("="*60)
print("ANÁLISIS COMPLETO DEL FORMATO Resumen_Dia.xlsx")
print("="*60)

print("\n=== DIMENSIONES ===")
print(f"Filas: {ws.max_row}, Columnas: {ws.max_column}")

print("\n=== IMÁGENES/LOGO ===")
if hasattr(ws, '_images') and ws._images:
    for idx, img in enumerate(ws._images):
        print(f"Imagen {idx + 1}:")
        print(f"  Ancla: {img.anchor}")
        print(f"  Width: {img.width}, Height: {img.height}")
else:
    print("No se encontraron imágenes embebidas")

print("\n=== CELDAS COMBINADAS ===")
for merged_range in ws.merged_cells:
    print(merged_range)

print("\n=== ESTRUCTURA DEL ENCABEZADO ===")
print("\nCELDA B2 (donde va el logo):")
cell = ws['B2']
print(f"  Valor: {cell.value}")
print(f"  Font: size={cell.font.sz}, bold={cell.font.b}")
print(f"  Fill: {cell.fill.patternType}")

print("\nCELDA C2 (título principal combinada C2:T2):")
cell = ws['C2']
print(f"  Valor: {cell.value}")
print(f"  Font: size={cell.font.sz}, bold={cell.font.b}, name={cell.font.name}")
print(f"  Color: {cell.font.color.rgb if cell.font.color else 'N/A'}")
print(f"  Fill: {cell.fill.fgColor.rgb if cell.fill.fgColor else 'N/A'}")
print(f"  Alignment: h={cell.alignment.horizontal}, v={cell.alignment.vertical}")

print("\nCELDA C3 (subtítulo combinada C3:T3):")
cell = ws['C3']
print(f"  Valor: {cell.value}")
print(f"  Font: size={cell.font.sz}, bold={cell.font.b}")
print(f"  Color: {cell.font.color.rgb if cell.font.color else 'N/A'}")

print("\nCELDA U2 (combinada U2:W2):")
cell = ws['U2']
print(f"  Valor: {cell.value}")
print(f"  Font: size={cell.font.sz}, bold={cell.font.b}")

print("\nCELDA U3 (combinada U3:W3):")
cell = ws['U3']
print(f"  Valor: {cell.value}")

print("\nFILA 5 (Mes):")
cell = ws['U5']
print(f"  U5 valor: {cell.value}")
print(f"  Font: size={cell.font.sz}, bold={cell.font.b}")

print("\n=== ANCHOS DE COLUMNAS (B-W) ===")
for col_letter in ['B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L', 'M', 
                   'N', 'O', 'P', 'Q', 'R', 'S', 'T', 'U', 'V', 'W']:
    width = ws.column_dimensions[col_letter].width
    print(f"Columna {col_letter}: {width}")

print("\n=== ALTOS DE FILAS (1-10) ===")
for row_num in range(1, 11):
    height = ws.row_dimensions[row_num].height
    print(f"Fila {row_num}: {height}")

print("\n=== FILA 6 (Encabezados de columnas) ===")
for col in range(2, 24):  # B a W
    cell = ws.cell(row=6, column=col)
    letra = get_column_letter(col)
    print(f"{letra}6: {cell.value}")
    if col == 2:  # Solo mostrar detalles para la primera
        print(f"    Font: size={cell.font.sz}, bold={cell.font.b}")
        print(f"    Fill: {cell.fill.fgColor.rgb if cell.fill.fgColor else 'N/A'}")
        print(f"    Borders: top={cell.border.top.style if cell.border.top else None}")

print("\n=== FILA 7 (Primera fila de datos) ===")
cell = ws['B7']
print(f"B7 (Fecha): {cell.value}")
print(f"  Formato: {cell.number_format}")
print(f"  Alignment: h={cell.alignment.horizontal}, v={cell.alignment.vertical}")

print("\n=== FILA FINAL (Totales) ===")
last_row = ws.max_row
cell = ws.cell(row=last_row, column=2)
print(f"Fila {last_row} - B: {cell.value}")
print(f"  Font: size={cell.font.sz}, bold={cell.font.b}")
print(f"  Color: {cell.font.color.rgb if cell.font.color else 'N/A'}")
print(f"  Fill: {cell.fill.fgColor.rgb if cell.fill.fgColor else 'N/A'}")
print(f"  Height: {ws.row_dimensions[last_row].height}")

print("\n" + "="*60)
print("ANÁLISIS COMPLETO")
print("="*60)
